"""Tests for excel_report.py and word_report.py, including traceability."""

from __future__ import annotations

import io

from docx import Document
from openpyxl import load_workbook

from src.excel_report import HEADERS, build_excel_report
from src.validator import validate_records
from src.word_report import build_word_report


def test_excel_has_headers_and_one_row_per_requirement(records):
    data = build_excel_report(records, "sample.pdf", page_count=2)
    workbook = load_workbook(io.BytesIO(data))

    sheet = workbook["Requirements"]
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet.max_row == len(records) + 1
    assert sheet.freeze_panes == "A2"
    assert "Summary" in workbook.sheetnames


def test_excel_keeps_the_source_page(records):
    data = build_excel_report(records, "sample.pdf", page_count=2)
    sheet = load_workbook(io.BytesIO(data))["Requirements"]

    page_column = HEADERS.index("Source Page") + 1
    pages = [sheet.cell(row=r, column=page_column).value for r in range(2, sheet.max_row + 1)]
    assert pages == ["1", "1", "2"]


def test_excel_marks_records_needing_review():
    from src.records import make_record

    record = make_record(
        {
            "section": "Security",
            "sub_section": "Access Control",
            "feature": "Role Based Access",
            "requirement": "The system shall restrict administrative functions.",
            "source_pages": [],
        },
        "REQ-0001",
    )
    validated, _ = validate_records([record], [1, 2])

    data = build_excel_report(validated, "sample.pdf", page_count=2)
    sheet = load_workbook(io.BytesIO(data))["Requirements"]

    review_column = HEADERS.index("Review Required") + 1
    assert sheet.cell(row=2, column=review_column).value == "YES"
    assert sheet.cell(row=2, column=HEADERS.index("Source Page") + 1).value == "NOT FOUND"


def test_word_contains_hierarchy_headings_and_pages(records):
    data = build_word_report(records, "sample.pdf", page_count=2)
    document = Document(io.BytesIO(data))

    text = "\n".join(p.text for p in document.paragraphs)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]

    assert "RFP Use Case Extraction Report" in text
    assert "Customer Management" in headings       # Section
    assert "Account Information" in headings       # Sub-Section
    assert "Balance Enquiry" in headings           # Feature
    assert "Source Page: 1" in text                # traceability


def test_word_handles_an_empty_requirement_list():
    data = build_word_report([], "sample.pdf", page_count=2)
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)

    assert "No requirements were extracted" in text
