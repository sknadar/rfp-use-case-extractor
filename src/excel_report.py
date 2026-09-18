"""
excel_report.py
===============

Single responsibility: write the .xlsx report with openpyxl.

Two worksheets:
* "Requirements" - one row per requirement, including the source page.
* "Summary"      - document name, counts and review statistics.

Nothing in this file knows about Gemini, PDFs or Streamlit, so it can be tested
with a handful of dictionaries.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.logging_config import get_logger
from src.records import format_pages

logger = get_logger(__name__)

HEADERS = [
    "Requirement ID",
    "Section",
    "Sub-Section",
    "Feature",
    "Requirement",
    "Source Page",
    "Business Context",
    "Confidence",
    "Review Required",
    "Review Reasons",
    "Modified By Reviewer",
    "Extracted At",
]

# Column widths in Excel "character" units, in the same order as HEADERS.
COLUMN_WIDTHS = [16, 24, 24, 24, 70, 14, 45, 12, 16, 45, 18, 22]

HEADER_FILL = PatternFill("solid", fgColor="16233B")     # ink navy, matches the app
HEADER_FONT = Font(name="Arial", bold=True, color="FAF7F1", size=11)
BODY_FONT = Font(name="Arial", size=10)
REVIEW_FILL = PatternFill("solid", fgColor="F4E1DC")     # clay-tinted, matches the app
THIN = Side(style="thin", color="D9D2C2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class ExcelReportError(Exception):
    """Raised when the workbook cannot be produced."""


def build_excel_report(
    records: Sequence[Dict[str, Any]],
    document_name: str,
    page_count: int,
    extra_summary: Optional[Dict[str, Any]] = None,
) -> bytes:
    """
    Build the workbook and return it as bytes (ready for a download button).

    Returning bytes instead of writing a file keeps this module usable from
    Streamlit, from a script, or from a test, without any path handling.
    """
    try:
        workbook = Workbook()
        _write_requirements_sheet(workbook, records)
        _write_summary_sheet(workbook, records, document_name, page_count, extra_summary)

        stream = io.BytesIO()
        workbook.save(stream)
        logger.info("Excel generated | rows=%s", len(records))
        return stream.getvalue()
    except Exception as exc:
        logger.error("Excel generation failed: %s", exc, exc_info=True)
        raise ExcelReportError(
            "The Excel report could not be created. Please try again."
        ) from exc


def _write_requirements_sheet(workbook: Workbook, records: Sequence[Dict[str, Any]]) -> None:
    sheet = workbook.active
    sheet.title = "Requirements"

    # Header row
    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = BORDER
    sheet.row_dimensions[1].height = 28

    # Data rows
    for row_index, record in enumerate(records, start=2):
        review_required = bool(record.get("review_required"))
        values = [
            record.get("requirement_id", ""),
            record.get("section", ""),
            record.get("sub_section", ""),
            record.get("feature", ""),
            record.get("requirement", ""),
            format_pages(record.get("source_pages")),
            record.get("business_context", ""),
            record.get("confidence", ""),
            "YES" if review_required else "NO",
            "; ".join(record.get("review_reasons") or []),
            "YES" if record.get("modified_by_reviewer") else "NO",
            record.get("extracted_at", ""),
        ]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            if review_required:
                cell.fill = REVIEW_FILL

    # Widths, frozen header and filter make the sheet usable for a reviewer.
    for column_index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    sheet.freeze_panes = "A2"
    last_row = max(len(records) + 1, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last_row}"


def _write_summary_sheet(
    workbook: Workbook,
    records: Sequence[Dict[str, Any]],
    document_name: str,
    page_count: int,
    extra_summary: Optional[Dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet("Summary")

    sections = {r.get("section", "") for r in records}
    features = {(r.get("section", ""), r.get("sub_section", ""), r.get("feature", "")) for r in records}
    review_count = sum(1 for r in records if r.get("review_required"))
    traced = sum(1 for r in records if r.get("source_pages"))

    rows = [
        ("RFP Use Case Extraction - Summary", ""),
        ("", ""),
        ("Source document", document_name),
        ("Pages in document", page_count),
        ("Requirements extracted", len(records)),
        ("Sections", len(sections)),
        ("Features", len(features)),
        ("Requirements with a source page", traced),
        ("Requirements needing review", review_count),
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for key, value in (extra_summary or {}).items():
        rows.append((str(key), value))

    for row_index, (label, value) in enumerate(rows, start=1):
        label_cell = sheet.cell(row=row_index, column=1, value=label)
        sheet.cell(row=row_index, column=2, value=value).font = BODY_FONT
        label_cell.font = Font(name="Arial", bold=(row_index == 1), size=13 if row_index == 1 else 10)

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 60
